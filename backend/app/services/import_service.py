"""线索 Excel/CSV 导入（PRD §5.6，api-contract §8）。

preview 解析并校验上传文件、返回 preview token 与逐行错误；commit 用该 token 重新校验并
复用线索创建/匹配逻辑，批次幂等、部分成功、逐行结果持久化到 import_jobs/import_rows。
"""

import csv
import io
import secrets
import uuid
from datetime import datetime, timezone

from openpyxl import Workbook, load_workbook
from sqlalchemy import or_, select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.core.errors import APIError, not_found
from app.models.import_job import ImportJob, ImportRow
from app.models.lead import INTAKE_IMPORT, Lead
from app.models.lead_event import LeadEvent
from app.models.lead_need_type import LeadNeedType
from app.models.lead_source import LeadSourceCategory
from app.models.user import User
from app.repositories import lead_repository
from app.schemas.lead import LeadCreate
from app.services import company_service
from app.services.lead_service import _default_source_category, _validate_min_content

IMPORT_HEADERS: dict[str, str] = {
    "occurred_at": "获取时间",
    "contact_name_raw": "姓名",
    "phone_raw": "手机号",
    "email_raw": "邮箱",
    "wechat_raw": "微信号",
    "company_name_raw": "企业名称",
    "need_type": "需求类型",
    "inquiry_text": "需求描述",
    "notes": "备注",
    "source_category": "来源类型",
    "source_detail_text": "来源详情",
}

# preview token → 解析后的行（进程内短时存储，commit 后清除）
_previews: dict[str, list[dict]] = {}


def build_template_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "线索导入模板"
    ws.append(list(IMPORT_HEADERS.values()))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _parse_datetime(value: str | None) -> datetime | None:
    if not value:
        return None
    value = value.strip()
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d", "%Y/%m/%d"):
        try:
            dt = datetime.strptime(value, fmt)
            return dt.replace(tzinfo=timezone.utc)
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(value)
    except ValueError:
        return None


def _parse_xlsx(content: bytes) -> list[dict]:
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    header = next(rows_iter, None)
    if header is None:
        return []
    header_map = {str(h).strip(): field for field, h in IMPORT_HEADERS.items()}
    out = []
    for r in rows_iter:
        if all(v is None or str(v).strip() == "" for v in r):
            continue
        row: dict[str, str] = {}
        for idx, h in enumerate(header):
            field = header_map.get(str(h).strip())
            if field and idx < len(r):
                v = r[idx]
                row[field] = "" if v is None else str(v).strip()
        out.append(row)
    return out


def _decode(content: bytes) -> str:
    for enc in ("utf-8-sig", "gbk"):
        try:
            return content.decode(enc)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def _parse_csv(content: bytes) -> list[dict]:
    text = _decode(content)
    reader = csv.DictReader(io.StringIO(text))
    header_map = {str(k).strip(): field for field, k in IMPORT_HEADERS.items()}
    out = []
    for raw in reader:
        row: dict[str, str] = {}
        for k, v in raw.items():
            if k is None:
                continue
            field = header_map.get(k.strip())
            if field:
                row[field] = (v or "").strip()
        if any(row.values()):
            out.append(row)
    return out


def _parse(filename: str, content: bytes) -> list[dict]:
    name = filename.lower()
    if name.endswith(".xlsx"):
        return _parse_xlsx(content)
    if name.endswith(".csv"):
        return _parse_csv(content)
    raise APIError(422, "VALIDATION_ERROR", "仅支持 .xlsx / .csv 文件")


def _row_to_payload(db: Session, row: dict) -> tuple[LeadCreate | None, str | None]:
    need_type_id = None
    if row.get("need_type"):
        nt = db.scalar(
            select(LeadNeedType).where(LeadNeedType.name == row["need_type"])
        )
        if nt is None:
            return None, f"需求类型不存在：{row['need_type']}"
        need_type_id = nt.id

    source_category_id = None
    if row.get("source_category"):
        sc = db.scalar(
            select(LeadSourceCategory).where(
                or_(
                    LeadSourceCategory.name == row["source_category"],
                    LeadSourceCategory.code == row["source_category"],
                )
            )
        )
        if sc is None:
            return None, f"来源类型不存在：{row['source_category']}"
        source_category_id = sc.id
    else:
        # 未填来源时归入「未知」，与手动创建一致；leads.source_category_id 为 NOT NULL
        source_category_id = _default_source_category(db).id

    occurred_at = _parse_datetime(row.get("occurred_at"))
    if row.get("occurred_at") and occurred_at is None:
        return None, f"获取时间格式错误：{row['occurred_at']}"

    payload = LeadCreate(
        occurred_at=occurred_at,
        contact_name_raw=row.get("contact_name_raw") or None,
        phone_raw=row.get("phone_raw") or None,
        email_raw=row.get("email_raw") or None,
        wechat_raw=row.get("wechat_raw") or None,
        company_name_raw=row.get("company_name_raw") or None,
        need_type_id=need_type_id,
        inquiry_text=row.get("inquiry_text") or None,
        notes=row.get("notes") or None,
        source_category_id=source_category_id,
        source_detail_text=row.get("source_detail_text") or None,
        intake_method=INTAKE_IMPORT,
    )
    try:
        _validate_min_content(payload)
    except APIError as e:
        return None, e.message
    return payload, None


def preview(db: Session, filename: str, content: bytes) -> dict:
    if len(content) > settings.import_max_file_size_mb * 1024 * 1024:
        raise APIError(422, "VALIDATION_ERROR", "文件超过大小上限")

    rows = _parse(filename, content)
    if len(rows) > settings.import_max_rows:
        raise APIError(
            422, "VALIDATION_ERROR", f"行数超过上限 {settings.import_max_rows}"
        )

    errors: list[dict] = []
    valid_rows = 0
    duplicates = 0
    for idx, row in enumerate(rows, start=1):
        payload, err = _row_to_payload(db, row)
        if err is not None:
            errors.append({"row_number": idx, "message": err})
            continue
        valid_rows += 1
        if payload.company_name_raw:
            _, candidates = company_service.match_company(db, payload.company_name_raw)
            if candidates:
                duplicates += 1

    token = secrets.token_urlsafe(16)
    _previews[token] = rows
    return {
        "preview_token": token,
        "total_rows": len(rows),
        "valid_rows": valid_rows,
        "duplicates": duplicates,
        "errors": errors,
    }


def _create_lead_row(db: Session, actor: User, payload: LeadCreate) -> Lead:
    lead = Lead(
        company_name_raw=payload.company_name_raw,
        contact_name_raw=payload.contact_name_raw,
        phone_raw=payload.phone_raw,
        email_raw=payload.email_raw,
        wechat_raw=payload.wechat_raw,
        need_type_id=payload.need_type_id,
        inquiry_text=payload.inquiry_text,
        notes=payload.notes,
        occurred_at=payload.occurred_at or datetime.now(timezone.utc),
        source_category_id=payload.source_category_id,
        source_detail_id=payload.source_detail_id,
        source_detail_text=payload.source_detail_text,
        intake_method=INTAKE_IMPORT,
        created_by=actor.id,
    )
    lead_repository.create_lead(db, lead)
    db.add(LeadEvent(lead_id=lead.id, event_type="lead.created", created_by=actor.id))
    matched, _ = company_service.match_company(db, payload.company_name_raw)
    if matched is not None:
        lead.company_id = matched.id
    return lead


def commit(db: Session, actor: User, preview_token: str, idempotency_key: str) -> dict:
    existing = db.scalar(
        select(ImportJob).where(ImportJob.idempotency_key == idempotency_key)
    )
    if existing is not None:
        return {
            "batch_id": str(existing.id),
            "success_rows": existing.success_rows,
            "failed_rows": existing.failed_rows,
            "skipped_rows": 0,
        }

    rows = _previews.pop(preview_token, None)
    if rows is None:
        raise not_found("preview token 无效或已过期")

    job = ImportJob(
        created_by=actor.id,
        filename="import",
        status="done",
        idempotency_key=idempotency_key,
        total_rows=len(rows),
    )
    db.add(job)
    db.flush()

    success = failed = 0
    for idx, row in enumerate(rows, start=1):
        payload, err = _row_to_payload(db, row)
        if err is not None:
            db.add(
                ImportRow(job_id=job.id, row_number=idx, status="failed", message=err)
            )
            failed += 1
            continue
        lead = _create_lead_row(db, actor, payload)
        db.add(
            ImportRow(job_id=job.id, row_number=idx, status="success", lead_id=lead.id)
        )
        success += 1

    job.success_rows = success
    job.failed_rows = failed
    db.commit()
    return {
        "batch_id": str(job.id),
        "success_rows": success,
        "failed_rows": failed,
        "skipped_rows": 0,
    }


def get_result(db: Session, job_id: uuid.UUID) -> dict:
    job = db.get(ImportJob, job_id)
    if job is None:
        raise not_found("导入批次不存在")
    rows = list(
        db.scalars(
            select(ImportRow)
            .where(ImportRow.job_id == job_id)
            .order_by(ImportRow.row_number)
        ).all()
    )
    return {
        "batch_id": str(job.id),
        "filename": job.filename,
        "status": job.status,
        "total_rows": job.total_rows,
        "success_rows": job.success_rows,
        "failed_rows": job.failed_rows,
        "rows": [
            {
                "row_number": r.row_number,
                "status": r.status,
                "message": r.message,
                "lead_id": str(r.lead_id) if r.lead_id else None,
            }
            for r in rows
        ],
    }
